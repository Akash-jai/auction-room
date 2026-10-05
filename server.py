"""Auction Room server: FastAPI + WebSockets. The server owns all auction rules, so clients can't cheat."""
import asyncio, base64, io, json, os, random, string

from fastapi import FastAPI, WebSocket, WebSocketDisconnect
from fastapi.responses import FileResponse, Response

from players import build_pool, order, pick
from custom import parse_upload
from tourney import simulate

PURSE, MAXS, MINS, MAXO, MAXO_XI = 120.0, 25, 18, 8, 4
# bidding clock in seconds: (first clock, after a bid) for big players (base >= 1 cr) and small ones
SPEEDS = {"quick": ((12, 8), (6, 4)), "normal": ((20, 12), (12, 8)), "relaxed": ((40, 25), (25, 15))}
TM = [("Mumbai", "#1c5fd1"), ("Chennai", "#f2b705"), ("Bengaluru", "#d3222a"), ("Kolkata", "#5b2a86"),
      ("Delhi", "#2a7fd4"), ("Hyderabad", "#f26a1b"), ("Punjab", "#c8102e"), ("Rajasthan", "#e8388f"),
      ("Gujarat", "#1b3a57"), ("Lucknow", "#22a5b8")]
GOAL = {"W": 2, "B": 6, "A": 5, "L": 8}


def r2(x): return round(x + 1e-9, 2)
def inc(x): return .1 if x < 1 else .25 if x < 2 else .5 if x < 5 else 1.0
def val(r): return max(.2, .2 + (r - 55) / 15) if r < 70 else 1.2 + ((r - 70) / 25) ** 2 * 15


class Team:
    def __init__(s, i, n, c, owner=None):
        s.i, s.n, s.c, s.owner = i, n, c, owner
        s.purse, s.sq = PURSE, []
        s.ag, s.st, s.sn = random.uniform(.85, 1.25), random.uniform(.8, 1.4), random.random() < .3
        s.target = random.randint(19, 24)  # AI stops buying around here, like real franchises
        s.xi, s.locked, s.final = [], False, None  # playing XI: picked names, locked flag, final players

    def cnt(s, k): return sum(1 for x in s.sq if x["k"] == k)
    def ov(s): return sum(1 for x in s.sq if x["o"])


RULES = (("W", 1, "wicketkeeper"), ("BW", 4, "batters (keeper counts)"), ("L", 3, "specialist bowlers"),
         ("LA", 5, "bowlers or all-rounders"))  # a proper XI: keeper, batting depth, bowling attack


def needs(sq):
    """Rule minimums, relaxed only if the squad simply doesn't contain enough players of that kind."""
    return [(ks, min(n, sum(1 for x in sq if x["k"] in ks)), label) for ks, n, label in RULES]


def xi_problems(sq, names):
    by = {p["n"]: p for p in sq}
    xi = [by[n] for n in names if n in by]
    return [f"at least {n} {label}" for ks, n, label in needs(sq) if sum(1 for x in xi if x["k"] in ks) < n]


def build_xi(sq, seed=()):
    """Best balanced XI. AI franchises use this; humans only if they never locked in. `seed` = players to keep."""
    by = {p["n"]: p for p in sq}
    sel = [by[n] for n in seed if n in by][:11]
    ps = sorted(sq, key=lambda x: -x["r"])

    def ok(x): return x not in sel and len(sel) < 11 and not (x["o"] and sum(1 for y in sel if y["o"]) >= MAXO_XI)
    def cnt(ks): return sum(1 for y in sel if y["k"] in ks)
    for ks, need in (("W", 1), ("L", 3), ("BW", 4), ("LA", 5)):
        for x in ps:
            if cnt(ks) >= need or len(sel) >= 11: break
            if x["k"] in ks and ok(x): sel.append(x)
    for x in ps:
        if len(sel) >= 11: break
        if ok(x): sel.append(x)
    return sel


def breakdown(t):
    xi = t.final or build_xi(t.sq)
    base = sum(p["r"] for p in xi)
    kb = 8 if any(p["k"] == "W" for p in xi) else 0
    bb = 10 if sum(1 for p in xi if p["k"] in "LA") >= 5 else 0
    tb = 6 if sum(1 for p in xi if p["k"] in "BWA") >= 5 else 0
    pen = -100 if len(t.sq) < MINS else 0
    return {"xi": xi, "base": base, "kb": kb, "bb": bb, "tb": tb, "pen": pen, "score": base + kb + bb + tb + pen}


def score(t):
    b = breakdown(t)
    return b["score"], b["xi"]


class Room:
    def __init__(s, code):
        s.code, s.cl, s.host, s.claims = code, {}, None, {}
        s.phase, s.teams, s.task = "lobby", [], None
        s.q, s.idx, s.p, s.cur, s.lead = [], -1, None, None, None
        s.tl = s.na = s.pause = s.t1 = 0
        s.status, s.last, s.log, s.unsold, s.pass2 = "bidding", "", [], [], False
        s.reaper = None
        s.bids = []  # bid history of the current lot, oldest first
        s.mode, s.speed, s.tour = "match", "normal", None  # match = one XI scores; tournament = depth wins
        s.custom = None  # host-uploaded player list: {"players": [...], "file": str, "notes": [...]}

    # ---------- connections ----------
    def attach(s, cid, name, ws):
        if cid in s.cl:
            s.cl[cid]["ws"] = ws
            if name: s.cl[cid]["name"] = name
        else:
            s.cl[cid] = {"name": (name or "Guest")[:16], "ws": ws}
        if s.host is None or not s.cl.get(s.host, {}).get("ws"): s.host = cid

    def detach(s, cid, ws):
        c = s.cl.get(cid)
        if c and c["ws"] is ws:
            c["ws"] = None
            if s.host == cid:
                live = [k for k, v in s.cl.items() if v["ws"]]
                if live: s.host = live[0]
        if s.phase == "lobby":  # free the claim of someone who left the lobby
            for i, o in list(s.claims.items()):
                if o == cid: del s.claims[i]

    def alive(s): return any(c["ws"] for c in s.cl.values())

    def auto(s, t): return not t.owner or not s.cl.get(t.owner, {}).get("ws")

    # ---------- rules ----------
    def nb(s): return s.p["b"] if s.cur is None else r2(s.cur + inc(s.cur))

    def can(s, t, p, a):
        keep = max(0, MINS - len(t.sq) - 1) * .25
        return len(t.sq) < MAXS and a <= t.purse - keep + 1e-9 and (not p["o"] or t.ov() < MAXO)

    def maxpay(s, t, p):
        if len(t.sq) >= (t.target if len(t.sq) >= MINS else MAXS): return 0
        c = t.cnt(p["k"])
        rem = sum(1 for x in s.q[s.idx:] if x["k"] == p["k"] and x["r"] >= p["r"] - 3)
        scar = 1.25 if rem <= 2 else .85 if rem >= 8 else 1
        nd = 1.35 if c < GOAL[p["k"]] else .5 if c >= GOAL[p["k"]] + 2 else .95
        if p["k"] == "W" and c >= 2: nd = .3
        m = t.ag * nd * scar * (t.st if p["r"] >= 88 else 1) * (1.2 if len(t.sq) < MINS else 1)
        left = max(.15, 1 - s.idx / max(1, len(s.q)))
        pressure = (t.purse / PURSE) / left
        if pressure > 1.15: m *= min(1.7, pressure)  # sitting on cash late: spend it
        v = val(p["r"]) * (1 + (m - 1) * .55)
        keep = max(0, MINS - len(t.sq) - 1) * .25
        return min(v, (t.purse - keep) * (.35 if p["r"] >= 88 else .25), t.purse - keep)

    # ---------- flow ----------
    def setup(s, size):
        humans = sorted(s.claims)
        size = max(2, min(10, size, 10)); size = max(size, len(humans))
        if s.custom and len(s.custom["players"]) < MINS * size:
            return (f"Your list has {len(s.custom['players'])} players, but {size} teams need at least {MINS * size} "
                    f"({MINS} each). Add players or pick fewer teams.")
        free = [i for i in range(10) if i not in s.claims]
        ais = random.sample(free, size - len(humans))
        s.teams = [Team(k, TM[i][0], TM[i][1], s.claims.get(i)) for k, i in enumerate(sorted(humans + ais))]
        if s.custom:  # only the uploaded players are auctioned
            pool = [dict(p) for p in s.custom["players"]]
            mq = max(5, min(30, round(len(pool) * .12)))
            q = order(pool, mq)
        else:
            pool = sorted(build_pool(), key=lambda p: -p["r"])
            n = min(len(pool), max(70, size * 26))  # fewer teams -> shorter auction
            mq = min(30, 14 + size * 2)  # marquee set grows with the number of teams (2 -> 18, 10 -> 30)
            q = order(pick(pool, n, mq), mq)
        s.q, s.idx, s.unsold, s.pass2, s.log = q, -1, [], False, []
        s.phase = "auction"
        s.next_lot()

    def next_lot(s):
        while True:
            s.idx += 1
            if all(len(t.sq) >= MAXS or (s.auto(t) and len(t.sq) >= t.target) for t in s.teams): return s.finish()
            if s.idx >= len(s.q):
                if not s.pass2 and s.unsold and any(len(t.sq) < MINS for t in s.teams):
                    s.pass2 = True
                    random.shuffle(s.unsold)
                    s.q += s.unsold; s.unsold = []
                    s.log.append("Accelerated round: unsold players return")
                    s.idx -= 1
                    continue
                return s.finish()
            p = s.q[s.idx]
            if not any(s.can(t, p, p["b"]) for t in s.teams):
                s.unsold.append(p); continue
            break
        (f1, a1), (f2, a2) = SPEEDS[s.speed]
        big = p["b"] >= 1
        s.t1 = a1 if big else a2
        s.p, s.cur, s.lead, s.status, s.bids = p, None, None, "bidding", []
        s.tl, s.na = (f1 if big else f2), random.randint(1, 3)

    def team_of(s, cid): return next((x for x in s.teams if x.owner == cid), None)

    def finish(s):  # auction over -> everyone picks a playing XI
        s.phase = "xi"
        for t in s.teams:
            t.final = None
            t.xi = [p["n"] for p in build_xi(t.sq)] if not t.owner else []  # humans start with an empty XI
            t.locked = not t.owner  # AI franchises lock instantly

    def set_xi(s, cid, names):
        t = s.team_of(cid)
        if not t or s.phase != "xi" or not isinstance(names, list): return None
        names = [n for n in names if isinstance(n, str)]
        by = {p["n"]: p for p in t.sq}
        if len(set(names)) != len(names) or any(n not in by for n in names): return "Invalid pick."
        if len(names) > 11: return "An XI has 11 players."
        if sum(1 for n in names if by[n]["o"]) > MAXO_XI: return f"Max {MAXO_XI} overseas players in the XI."
        t.xi, t.locked = names, False
        return None

    def lock(s, cid):
        t = s.team_of(cid)
        if not t or s.phase != "xi": return None
        if len(t.xi) != min(11, len(t.sq)): return "Pick 11 players first."
        bad = xi_problems(t.sq, t.xi)
        if bad: return "Not a proper XI. You need " + ", ".join(bad) + "."
        t.locked = True
        if all(x.locked for x in s.teams): s.reveal()
        return None

    def reveal(s):
        for t in s.teams:
            t.final = build_xi(t.sq, t.xi)
        s.tour = simulate(s.teams, build_xi, lambda t: s.cl.get(t.owner, {}).get("name")) if s.mode == "tournament" else None
        s.phase = "end"

    def place(s, t, a):
        s.cur, s.lead, s.tl, s.na = a, t, s.t1, random.randint(1, 3)
        s.bids.append({"n": t.n, "c": t.c, "a": a})

    def bid(s, cid):
        t = next((x for x in s.teams if x.owner == cid), None)
        if s.phase != "auction" or s.status != "bidding" or not t: return
        a = s.nb()
        if t is not s.lead and s.can(t, s.p, a): s.place(t, a)

    def hammer(s):
        s.status, s.pause = "result", 3
        if s.lead:
            s.lead.purse = r2(s.lead.purse - s.cur)
            s.lead.sq.append({**s.p, "paid": s.cur})
            s.last = f"SOLD to {s.lead.n} for ₹{s.cur} cr"
            s.log.append(f"{s.p['n']} → {s.lead.n}, ₹{s.cur} cr")
        else:
            s.unsold.append(s.p); s.last = "UNSOLD"
            s.log.append(f"{s.p['n']}: unsold")

    def tick(s):
        if s.phase != "auction": return
        if s.status == "result":
            s.pause -= 1
            if s.pause <= 0: s.next_lot()
            return
        s.tl -= 1; s.na -= 1
        a, p = s.nb(), s.p
        hot = s.lead is not None and not s.auto(s.lead)
        c = [t for t in s.teams if s.auto(t) and t is not s.lead and s.can(t, p, a) and a <= s.maxpay(t, p)
             and (not t.sn or s.tl <= 8) and s.na <= 0 and random.random() < (.85 if hot else .7)]
        if c: s.place(random.choice(c), a)
        if s.tl <= 0: s.hammer()

    async def run(s):
        try:
            while s.phase == "auction":
                await asyncio.sleep(1)
                if not s.alive(): continue  # everyone left (e.g. page reload): freeze the clock
                s.tick()
                await s.broadcast()
        except asyncio.CancelledError:
            pass

    # ---------- messages ----------
    async def handle(s, cid, m):
        t, c = m.get("t"), s.cl[cid]
        if t == "pick" and s.phase == "lobby":
            i = m.get("i")
            if not isinstance(i, int) or not 0 <= i < 10: return
            if s.claims.get(i) == cid: del s.claims[i]
            elif i not in s.claims:
                for k, o in list(s.claims.items()):
                    if o == cid: del s.claims[k]
                s.claims[i] = cid
        elif t == "start" and cid == s.host and s.phase == "lobby":
            if not s.claims:
                if c["ws"]: await c["ws"].send_text(json.dumps({"t": "err", "m": "Pick a franchise first."}))
                return
            err = s.setup(int(m.get("size", 10)))
            if err:
                if c["ws"]: await c["ws"].send_text(json.dumps({"t": "err", "m": err}))
                return
            if s.phase == "auction": s.task = asyncio.create_task(s.run())
        elif t == "bid":
            s.bid(cid)
        elif t == "end" and cid == s.host and s.phase == "auction":
            s.finish()
        elif t == "upload" and cid == s.host and s.phase == "lobby":
            try:
                data = base64.b64decode(m.get("data") or "", validate=False)
                players, notes = parse_upload(str(m.get("name") or "")[:80], data)
                s.custom = {"players": players, "file": str(m.get("name") or "")[:60], "notes": notes}
            except ValueError as e:
                if c["ws"]: await c["ws"].send_text(json.dumps({"t": "err", "m": str(e)}))
        elif t == "opts" and cid == s.host and s.phase == "lobby":
            if m.get("mode") in ("match", "tournament"): s.mode = m["mode"]
            if m.get("speed") in SPEEDS: s.speed = m["speed"]
        elif t == "nocustom" and cid == s.host and s.phase == "lobby":
            s.custom = None
        elif t in ("xi", "lock", "unlock", "reveal") and s.phase == "xi":
            me, err = s.team_of(cid), None
            if t == "xi": err = s.set_xi(cid, m.get("names"))
            elif t == "lock": err = s.lock(cid)
            elif t == "unlock" and me: me.locked = False
            elif t == "reveal" and cid == s.host: s.reveal()
            if err and c["ws"]: await c["ws"].send_text(json.dumps({"t": "err", "m": err}))
        elif t == "lobby" and cid == s.host and s.phase == "end":
            s.phase, s.teams, s.claims, s.tour = "lobby", [], {}, None
        await s.broadcast()

    def state(s, cid):
        me = next((t for t in s.teams if t.owner == cid), None)
        d = {"t": "state", "code": s.code, "phase": s.phase,
             "you": {"host": cid == s.host, "team": me.i if me else None,
                     "claim": next((i for i, o in s.claims.items() if o == cid), None)},
             "people": [{"n": c["name"], "on": bool(c["ws"]), "host": k == s.host} for k, c in s.cl.items()],
             "fr": [{"n": n, "c": c, "o": s.cl[s.claims[i]]["name"] if s.claims.get(i) in s.cl else None}
                    for i, (n, c) in enumerate(TM)],
             "limits": {"purse": PURSE, "min": MINS, "max": MAXS, "ov": MAXO}, "custom": None,
             "mode": s.mode, "speed": s.speed}
        if s.custom:
            ps = s.custom["players"]
            d["custom"] = {"file": s.custom["file"], "n": len(ps), "notes": s.custom["notes"], "ov": sum(p["o"] for p in ps),
                           **{k: sum(1 for p in ps if p["k"] == k) for k in "BLAW"}}
        if s.phase != "lobby":
            d["teams"] = [{"i": t.i, "n": t.n, "c": t.c, "o": s.cl.get(t.owner, {}).get("name"), "purse": r2(t.purse),
                           "cnt": len(t.sq), "ov": t.ov(), **{k: t.cnt(k) for k in "BLAW"}} for t in s.teams]
        if s.phase == "auction":
            a = s.nb()
            d["lot"] = {"i": s.idx + 1, "of": len(s.q), "p": s.p, "cur": s.cur, "lead": s.lead.i if s.lead else None,
                        "tl": s.tl, "st": s.status, "nb": a, "last": s.last,
                        "bids": s.bids[-8:][::-1], "nbids": len(s.bids),
                        "can": bool(me and s.status == "bidding" and me is not s.lead and s.can(me, s.p, a))}
            d["mine"] = me.sq if me else []
            d["log"] = s.log[-10:][::-1]
        if s.phase == "xi":
            d["mine"] = me.sq if me else []
            d["xi"] = {"sel": me.xi if me else [], "locked": bool(me and me.locked), "max_ov": MAXO_XI,
                       "rules": [{"ks": ks, "need": n, "label": label} for ks, n, label in needs(me.sq)] if me else [],
                       "wait": [{"n": t.n, "who": s.cl.get(t.owner, {}).get("name"), "ok": t.locked}
                                for t in s.teams if t.owner]}
        if s.phase == "end":
            res = []
            for t in s.teams:
                b = breakdown(t)
                res.append({"n": t.n, "c": t.c, "i": t.i, "o": s.cl.get(t.owner, {}).get("name"), "score": b["score"],
                            "base": b["base"], "kb": b["kb"], "bb": b["bb"], "tb": b["tb"], "pen": b["pen"],
                            "purse": r2(t.purse), "xi": b["xi"], "sq": t.sq})
            d["res"] = sorted(res, key=lambda x: -x["score"])
            d["tour"] = s.tour
        return d

    async def broadcast(s):
        for cid, c in list(s.cl.items()):
            if c["ws"]:
                try: await c["ws"].send_text(json.dumps(s.state(cid)))
                except Exception: c["ws"] = None


app = FastAPI()
rooms = {}


@app.get("/template.xlsx")
def template():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = "Players"
    ws.append(["Name", "Rating", "Position", "Overseas"])
    for r in [("Virat Kohli", 92, "Batter", "No"), ("Jasprit Bumrah", 95, "Bowler", "No"), ("Rishabh Pant", 90, "Keeper", "No"),
              ("Hardik Pandya", 90, "All-rounder", "No"), ("Jos Buttler", 90, "Keeper", "Yes"), ("Rashid Khan", 92, "Bowler", "Yes"),
              ("Andre Russell", 85, "All-rounder", "Yes"), ("Travis Head", 90, "Batter", "Yes")]: ws.append(r)
    for c in ws[1]: c.font = Font(bold=True)
    for col, w in zip("ABCD", (24, 10, 14, 11)): ws.column_dimensions[col].width = w
    h = wb.create_sheet("Help")
    for line in ["Only the first sheet (Players) is read.", "Name: player name (must be unique).",
                 "Rating: number, 60 to 95 works best (higher = better = pricier).",
                 "Position: Batter, Bowler, All-rounder or Keeper.", "Overseas: Yes or No.",
                 "Replace the example rows with your own players. You need at least 18 players per team."]: h.append([line])
    h.column_dimensions["A"].width = 80
    b = io.BytesIO(); wb.save(b)
    return Response(b.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="auction-players-template.xlsx"'})


async def reap(code, room, grace=600):
    try:
        await asyncio.sleep(grace)
    except asyncio.CancelledError:
        return
    if not room.alive():
        if room.task: room.task.cancel()
        if rooms.get(code) is room: rooms.pop(code, None)

HERE = os.path.dirname(os.path.abspath(__file__))


@app.get("/")
async def index(): return FileResponse(os.path.join(HERE, "static", "index.html"))


@app.get("/health")
async def health(): return {"ok": True, "rooms": len(rooms)}


@app.websocket("/ws/{code}")
async def ws_ep(ws: WebSocket, code: str, cid: str = "", name: str = ""):
    await ws.accept()
    code = "".join(ch for ch in code.upper() if ch.isalnum())[:8] or "MAIN"
    room = rooms.setdefault(code, Room(code))
    cid = (cid or "".join(random.choices(string.ascii_lowercase, k=10)))[:24]
    if room.reaper: room.reaper.cancel(); room.reaper = None
    room.attach(cid, name, ws)
    await room.broadcast()
    try:
        while True:
            m = json.loads(await ws.receive_text())
            if isinstance(m, dict):
                if m.get("t") == "name": room.cl[cid]["name"] = str(m.get("n", "Guest"))[:16]
                await room.handle(cid, m)
    except (WebSocketDisconnect, json.JSONDecodeError):
        pass
    finally:
        room.detach(cid, ws)
        if not room.alive():
            room.reaper = asyncio.create_task(reap(code, room))  # grace period so reloads don't kill the game
        else:
            await room.broadcast()
